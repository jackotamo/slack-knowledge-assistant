"""Real tiny documents generated locally: no network or API calls."""
import tempfile
import unittest
from pathlib import Path
from assistant.documents import DocumentCollection, MAX_UPLOAD_BYTES
from assistant.loader import parse_document
from assistant.search import search


def make_pdf(path):
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    writer = PdfWriter()
    for text in ("Incident reporting instructions", "Report service outage to coordinator"):
        page = writer.add_blank_page(width=300, height=300)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                 NameObject("/Subtype"): NameObject("/Type1"),
                                 NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 20 200 Td ({text}) Tj ET".encode("ascii"))
        page[NameObject("/Contents")] = writer._add_object(stream)
    writer.write(path)


class DocumentTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.collection = DocumentCollection(self.root / "uploads")

    def tearDown(self):
        self.temporary.cleanup()

    def test_markdown_metadata_and_management(self):
        content = b"# Example Policy\n\n## Reset password\nOpen the password portal.\n"
        path = self.collection.add("policy.md", content)
        passages, errors = self.collection.load()
        self.assertEqual(errors, [])
        self.assertEqual(passages[0].title, "Example Policy")
        self.assertEqual((passages[0].start_line, passages[0].end_line), (4, 4))
        self.assertEqual(search("reset password", passages)[0].passage.path, path.as_posix())
        self.assertEqual(self.collection.list_documents(), [path])
        self.collection.remove("policy.md")
        self.assertEqual(self.collection.load(), ([], []))
        with self.assertRaises(ValueError):
            self.collection.remove("policy.md")

    def test_invalid_uploads_leave_collection_unchanged(self):
        for name, data in (("../outside.md", b"text"), ("C:\\outside.md", b"text"),
                           ("CON.md", b"text"), ("script.py", b"text"),
                           ("empty.md", b""), ("blank.md", b" \n"),
                           ("broken.pdf", b"not a pdf"), ("broken.docx", b"not a zip"),
                           ("encoding.md", b"\xff")):
            with self.subTest(name=name), self.assertRaises(ValueError):
                self.collection.add(name, data)
        self.assertEqual(self.collection.list_documents(), [])

    def test_duplicate_does_not_overwrite(self):
        path = self.collection.add("policy.md", b"original text")
        with self.assertRaises(ValueError):
            self.collection.add("policy.md", b"different text")
        self.assertEqual(path.read_bytes(), b"original text")

    def test_oversized_upload_rejected(self):
        with self.assertRaises(ValueError):
            self.collection.add("large.md", b"x" * (MAX_UPLOAD_BYTES + 1))

    def test_pdf_pages_and_retrieval(self):
        path = self.root / "outage.pdf"
        make_pdf(path)
        passages = parse_document(path)
        self.assertEqual([p.page_number for p in passages], [1, 2])
        result = search("service outage coordinator", passages)[0]
        self.assertEqual(result.passage.page_number, 2)
        self.assertIn("coordinator", result.passage.text)
        self.collection.add(path.name, path.read_bytes())

    def test_blank_and_encrypted_pdf_rejected(self):
        from pypdf import PdfWriter
        path = self.root / "blank.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        writer.write(path)
        with self.assertRaisesRegex(ValueError, "No readable text"):
            parse_document(path)
        writer.encrypt("test-password")
        writer.write(path)
        with self.assertRaisesRegex(ValueError, "Encrypted"):
            parse_document(path)

    def test_docx_headings_tables_and_locations(self):
        from docx import Document
        path = self.root / "equipment.docx"
        document = Document()
        document.add_heading("Equipment procedure", level=1)
        document.add_paragraph("Collect your laptop from IT.")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Missing equipment"
        table.cell(0, 1).text = "Contact manager"
        document.save(path)
        passages = parse_document(path)
        self.assertEqual(passages[0].section, "Equipment procedure")
        self.assertEqual(passages[0].location_kind, "blocks")
        self.assertEqual((passages[0].start_line, passages[0].end_line), (2, 3))
        self.assertIn("Contact manager", passages[0].text)
        self.assertTrue(search("missing equipment", passages))
        self.collection.add(path.name, path.read_bytes())

    def test_corrupt_file_does_not_hide_valid_documents(self):
        self.collection.add("good.md", b"Useful procedure text")
        (self.collection.directory / "bad.docx").write_bytes(b"broken")
        passages, errors = self.collection.load()
        self.assertEqual(len(passages), 1)
        self.assertEqual(errors[0][0], "bad.docx")


if __name__ == "__main__":
    unittest.main()
